from fastapi import APIRouter,HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List,Dict,Any,Optional
from io import BytesIO
from app.services.bid_assembly_service import build_bid,list_bids
router=APIRouter(prefix="/bids",tags=["Bids"])
class In(BaseModel):analysis_results:List[Dict[str,Any]];company_name:Optional[str]=None;title:str="Bid Response"
@router.post("/assemble")
def assemble(x:In):return build_bid(x.analysis_results,x.company_name,x.title)
@router.get("")
def all():return list_bids()
def _find(id):
 for x in list_bids():
  if x["id"]==id:return x
 raise HTTPException(404,"Bid not found")
@router.get("/{id}/export/txt")
def txt(id:str):
 b=_find(id);s=b["title"]+"\n\n"
 for k,v in b["sections"].items():
  s+=k+"\n"+"-"*len(k)+"\n"
  for a in v:s+=a.get("requirement","")+": "+a.get("response","")+"\n\n"
 return StreamingResponse(BytesIO(s.encode()),media_type="text/plain",headers={"Content-Disposition":f'attachment; filename="bid_{id}.txt"'})
@router.get("/{id}/export/docx")
def docx(id:str):
 b=_find(id)
 try:
  from docx import Document
  d=Document();d.add_heading(b["title"],0)
  for k,v in b["sections"].items():
   d.add_heading(k,1)
   for a in v:d.add_paragraph(a.get("requirement",""),style="Heading 2" if a.get("requirement") else None);d.add_paragraph(a.get("response",""))
  bio=BytesIO();d.save(bio);bio.seek(0);return StreamingResponse(bio,media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",headers={"Content-Disposition":f'attachment; filename="bid_{id}.docx"'})
 except Exception:return txt(id)


@router.get("/{id}/export/pdf")
def pdf(id: str):
    b = _find(id)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch

        bio = BytesIO()
        doc = SimpleDocTemplate(
            bio, pagesize=A4,
            rightMargin=0.55*inch, leftMargin=0.55*inch,
            topMargin=0.55*inch, bottomMargin=0.55*inch
        )
        styles = getSampleStyleSheet()
        story = [Paragraph(b["title"], styles["Title"]), Spacer(1, 12)]

        for section, items in b.get("sections", {}).items():
            story.append(Paragraph(section, styles["Heading2"]))
            for item in items:
                req = item.get("requirement", "")
                response = item.get("response", "")
                if req:
                    story.append(Paragraph(req, styles["Heading4"]))
                story.append(Paragraph(response or "No response available.", styles["BodyText"]))
                story.append(Spacer(1, 8))

        doc.build(story)
        bio.seek(0)
        return StreamingResponse(
            bio,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="bid_{id}.pdf"'}
        )
    except Exception as e:
        raise HTTPException(500, f"PDF export failed: {e}")
